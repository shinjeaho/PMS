"""Exercise the export function in isolation without application or database setup."""
import ast
import io
from datetime import datetime
from pathlib import Path
import traceback
import unittest

from flask import Flask, jsonify, request, send_file
from openpyxl import load_workbook
import xlsxwriter


class AnnualMoneyExportTest(unittest.TestCase):
    def test_new_columns_and_four_balance_buckets(self):
        source = Path('pms/blueprints/PMS_annualProject.py').read_text(encoding='utf-8')
        tree = ast.parse(source)
        function = next(node for node in tree.body if isinstance(node, ast.FunctionDef) and node.name == 'export_annual_money')
        function.decorator_list = []
        namespace = dict(io=io, datetime=datetime, traceback=traceback, request=request,
                         send_file=send_file, jsonify=jsonify, xlsxwriter=xlsxwriter)
        exec(compile(ast.Module(body=[function], type_ignores=[]), '<export>', 'exec'), namespace)
        app = Flask(__name__)
        app.add_url_rule('/export', view_func=namespace['export_annual_money'], methods=['POST'])
        payload = {
            'year': 2026, 'vatMode': 'include',
            'stats': {'receipt': {'balanceNext': 1234}, 'pay': {'balanceNext': 5678}},
            'sections': [{'title': '차기년도 사업 목록', 'count': 1,
                          'rows': [{'no': 1, 'totalProgress': 42.5, 'completionBeforeTotal': 1100, 'completionTotal': 2200}],
                          'summary': {'completionBeforeTotal': 1100, 'completionTotal': 2200}},
                         {'title': '장기사업 목록', 'rows': []},
                         {'title': '차기년도 사업 목록', 'count': 1, 'groups': [
                             {'title': '차기년도(당해년도)', 'summaryLabel': '소계',
                              'rows': [{'contractCode': 'NEXT-1', 'receiptBalance': 123}],
                              'summary': {'receiptBalance': 123}},
                             {'title': '차기년도(장기사업)', 'summaryLabel': '소계', 'rows': []},
                         ]}],
        }
        response = app.test_client().post('/export', json=payload)
        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(io.BytesIO(response.data))
        stats, projects = workbook.worksheets
        self.assertEqual(stats['H7'].value, '차기년도')
        self.assertEqual(stats['H8'].value, 1234)
        self.assertEqual(stats['S8'].value, 5678)
        self.assertEqual(stats['L6'].value, '기지급')
        self.assertEqual(projects['D6'].value, '진행률')
        self.assertEqual(projects['D7'].value, '42.5%')
        self.assertEqual(projects['O6'].value, '준공금 기수령')
        self.assertEqual(projects['O7'].value, 1100)
        self.assertEqual(projects['P7'].value, 2200)
        self.assertEqual(projects['O8'].value, 1100)
        labels = {row[0].value: row[0].row for row in projects.iter_rows()
                  if row[0].value in ('차기년도(당해년도)', '차기년도(장기사업)')}
        current_row = labels['차기년도(당해년도)']
        long_row = labels['차기년도(장기사업)']
        self.assertEqual(projects.cell(current_row + 1, 2).value, 'NEXT-1')
        self.assertEqual(projects.cell(current_row + 2, 1).value, '소계')
        self.assertEqual(projects.cell(current_row + 2, 17).value, 123)
        self.assertEqual(long_row, current_row + 3)
        self.assertEqual(projects.cell(long_row + 2, 1).value, '소계')
        self.assertEqual(projects.cell(long_row + 2, 17).value, 0)


if __name__ == '__main__':
    unittest.main()
